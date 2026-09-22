"""Stage 02 - Requisition builder.

Takes the normalised table produced by stage 01, enriches every line from the
item master, works out the value SAP expects, and fills the requisition
workbook the stage 03 macro reads.

This stage knows the internal side of the process: the item master, the
tooling and service split, and the requisition layout. It knows nothing about
supplier file formats.

Usage:
    python build_requisition.py --input ../../01-supplier-intake/src/normalised_items.xlsx
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared"))
import parsers as P  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
MASTER_PATH = ROOT / "shared" / "master_data.xlsx"
TEMPLATE_PATH = ROOT / "shared" / "Input_RC_automatic.xlsx"

FIRST_DATA_ROW = 8
SHEET_BY_TYPE = {"MATERIAL": "Tooling", "SERVICE": "Service"}

# Column positions in the requisition template, identical on both sheets.
COL = {"code": 3, "quantity": 4, "value": 5, "description": 6,
       "item_type": 7, "material_group": 8, "budget_code": 9,
       "budget_line": 10, "project_code": 11}


def load_master(path=MASTER_PATH):
    """Item master keyed by the part number the supplier uses."""
    sheet = openpyxl.load_workbook(path, data_only=True)["Item Master"]
    headers = [c.value for c in sheet[4]]
    index = {name: i for i, name in enumerate(headers) if name}

    master = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        part = P.clean_text(row[index["Supplier Part Number"]])
        if not part:
            continue
        master[part.upper()] = {
            "internal_code": row[index["Internal Code"]],
            "item_category": row[index["Item Category"]],
            "material_group": row[index["Material Group"]],
            "budget_code": row[index["Budget Code"]],
            "budget_line": row[index["Budget Line"]],
            "project_code": row[index["Project Code"]],
            "operation_type": row[index["Operation Type"]],
        }
    return master


def load_normalised(path):
    sheet = openpyxl.load_workbook(path, data_only=True).active
    headers = [c.value for c in sheet[1]]
    return [dict(zip(headers, values))
            for values in sheet.iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in values)]


def sap_value(row):
    """Tooling takes the net amount plus IPI. Service takes it with all taxes.

    This is the one piece of business logic in the pipeline: the requisition
    does not carry tax rates, it carries the value those rates produce, and
    the rule differs by operation type.
    """
    rates = ("ipi",) if row["operation_type"] == "MATERIAL" else ("ipi", "icms", "pis_cofins", "iss")
    total_rate = sum(row.get(name) or 0 for name in rates)
    return round(row["net_amount"] * (1 + total_rate), 2)


def enrich(rows, master):
    """Apply the item master. A line whose item is not registered is rejected
    rather than guessed: a wrong material group creates a wrong requisition."""
    accepted, rejected = [], []

    for row in rows:
        part = P.clean_text(row.get("part_number"))
        entry = master.get(part.upper()) if part else None

        if not entry:
            row["problems"] = [f"'{part}' is not in the item master"]
            rejected.append(row)
            continue

        row.update(entry)
        row["sap_value"] = sap_value(row)
        accepted.append(row)

    return accepted, rejected


def write_requisition(rows, out_path, template=TEMPLATE_PATH):
    workbook = openpyxl.load_workbook(template)
    counters = {"Tooling": 0, "Service": 0}

    for row in rows:
        sheet_name = SHEET_BY_TYPE[row["operation_type"]]
        sheet = workbook[sheet_name]
        counters[sheet_name] += 1
        excel_row = FIRST_DATA_ROW + counters[sheet_name] - 1

        values = {
            "code": row["internal_code"],
            "quantity": row["quantity"],
            "value": row["sap_value"],
            "description": row["description"],
            "item_type": row["item_category"],
            "material_group": row["material_group"],
            "budget_code": row["budget_code"],
            "budget_line": row["budget_line"],
            "project_code": row["project_code"],
        }
        for field, column in COL.items():
            sheet.cell(row=excel_row, column=column, value=values[field])

    workbook.save(out_path)
    return counters


def write_report(accepted, rejected, counters, report_path):
    lines = [
        "Stage 02 - requisition builder",
        f"Run at {datetime.now():%Y-%m-%d %H:%M}",
        "",
        f"{len(accepted) + len(rejected)} normalised lines read",
        f"{len(accepted)} written ({counters['Tooling']} tooling, {counters['Service']} service)",
        f"{len(rejected)} rejected",
    ]
    if rejected:
        lines += ["", "Rejected lines:"]
        for row in rejected:
            where = f"{row.get('source_file')} row {row.get('source_row')}"
            lines.append(f"  {where}: {'; '.join(row['problems'])}")

    text = "\n".join(lines)
    Path(report_path).write_text(text, encoding="utf-8")
    return text


def main():
    parser = argparse.ArgumentParser(description="Stage 02 - requisition builder")
    parser.add_argument("--input", required=True, help="normalised table from stage 01")
    parser.add_argument("--output", default="Input_RC_automatic_filled.xlsx")
    parser.add_argument("--report", default="builder_report.txt")
    args = parser.parse_args()

    rows = load_normalised(Path(args.input))
    print(f"  {len(rows)} normalised lines read")

    accepted, rejected = enrich(rows, load_master())
    counters = write_requisition(accepted, args.output)
    print()
    print(write_report(accepted, rejected, counters, args.report))
    print(f"\nWritten to {args.output}")


if __name__ == "__main__":
    main()
